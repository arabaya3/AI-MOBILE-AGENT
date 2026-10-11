"""Advanced PowerPoint processing with slide context."""

from pathlib import Path
from typing import Any, Dict, List

from pptx import Presentation

from src.services.chunking.recursive import chunker


class AdvancedPowerPointProcessor:
    """Enhanced PowerPoint processor with slide and notes extraction."""

    async def process(self, file_path: Path) -> List[Dict[str, Any]]:
        """Process PowerPoint with slide-aware chunking."""
        chunks = []
        
        prs = Presentation(file_path)
        
        for slide_num, slide in enumerate(prs.slides, 1):
            slide_content = []
            
            title = self._extract_title(slide)
            if title:
                slide_content.append(f"Slide {slide_num}: {title}")
            else:
                slide_content.append(f"Slide {slide_num}")
            
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_content.append(shape.text.strip())
                
                if shape.has_table:
                    table_text = self._extract_table(shape.table)
                    if table_text:
                        slide_content.append(f"[TABLE]\n{table_text}")
            
            if slide.has_notes_slide:
                notes_text = slide.notes_slide.notes_text_frame.text
                if notes_text.strip():
                    slide_content.append(f"[NOTES]\n{notes_text.strip()}")
            
            full_text = "\n\n".join(slide_content)
            
            if full_text.strip():
                slide_chunks = await chunker.chunk_text(
                    full_text,
                    metadata={
                        "slide_number": slide_num,
                        "slide_title": title or f"Slide {slide_num}",
                        "total_slides": len(prs.slides),
                    }
                )
                chunks.extend(slide_chunks)
        
        return chunks

    def _extract_title(self, slide) -> str:
        """Extract slide title if present."""
        if slide.shapes.title:
            return slide.shapes.title.text.strip()
        return ""

    def _extract_table(self, table) -> str:
        """Extract table content as formatted text."""
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(" | ".join(cells))
        return "\n".join(rows)


powerpoint_processor = AdvancedPowerPointProcessor()
